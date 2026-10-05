from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
import math
import re


class FierroOrderParseError(RuntimeError):
    pass


@dataclass(frozen=True)
class FierroOrderLine:
    diametro_mm: float
    largo_m: float
    kg_mt: float
    barras_solicitadas: int
    kg_solicitados: float


@dataclass(frozen=True)
class FierroOrder:
    numero_pedido: str
    area: str
    fecha_pedido: date
    fecha_requerida_obra: date
    hora_requerida: str
    lineas: tuple[FierroOrderLine, ...]

    @property
    def total_kg(self) -> float:
        return round(sum(row.kg_solicitados for row in self.lineas), 2)


@dataclass
class _SheetMatrix:
    values: list[list[Any]]

    @property
    def max_row(self) -> int:
        return len(self.values)

    @property
    def max_col(self) -> int:
        return max((len(row) for row in self.values), default=0)

    def value(self, row: int, col: int) -> Any:
        if row < 0 or col < 0:
            return None
        if row >= len(self.values):
            return None
        current = self.values[row]
        if col >= len(current):
            return None
        return current[col]


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    text = (
        text.replace("Á", "A")
        .replace("É", "E")
        .replace("Í", "I")
        .replace("Ó", "O")
        .replace("Ú", "U")
        .replace("Ü", "U")
        .replace("Ñ", "N")
        .replace("º", "°")
    )
    text = re.sub(r"\s+", " ", text)
    return text


def _is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _as_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        if math.isfinite(number):
            return number
        return None

    text = str(value).strip().replace(" ", "")
    if not text:
        return None

    # Soporta 1.234,56 y 1234.56 sin asumir separador cuando no hace falta.
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")

    try:
        number = float(text)
    except ValueError:
        return None

    return number if math.isfinite(number) else None


def _as_int(value: Any) -> int | None:
    number = _as_number(value)
    if number is None:
        return None
    rounded = int(round(number))
    if abs(number - rounded) > 1e-6:
        return None
    return rounded


def _as_date(value: Any, label: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value or "").strip()
    formats = (
        "%m/%d/%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%d-%m-%Y",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    raise FierroOrderParseError(
        f"No se pudo interpretar {label}: {value!r}."
    )


def _as_time_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")

    if isinstance(value, (int, float)):
        # Fracción de día Excel.
        fraction = float(value) % 1
        total_minutes = int(round(fraction * 24 * 60)) % (24 * 60)
        hh, mm = divmod(total_minutes, 60)
        return f"{hh:02d}:{mm:02d}"

    return str(value).strip()


def _first_nonblank_right(sheet: _SheetMatrix, row: int, col: int) -> Any:
    for current_col in range(col + 1, sheet.max_col):
        value = sheet.value(row, current_col)
        if not _is_blank(value):
            return value
    return None


def _find_label_value(
    sheet: _SheetMatrix,
    aliases: tuple[str, ...],
    field_name: str,
) -> Any:
    normalized_aliases = tuple(_normalize_text(x) for x in aliases)

    for row in range(sheet.max_row):
        for col in range(sheet.max_col):
            current = _normalize_text(sheet.value(row, col))
            if not current:
                continue
            if any(alias in current for alias in normalized_aliases):
                value = _first_nonblank_right(sheet, row, col)
                if not _is_blank(value):
                    return value

    raise FierroOrderParseError(
        f"No se encontró el campo '{field_name}' en el archivo."
    )


def _find_detail_header(sheet: _SheetMatrix) -> tuple[int, dict[str, int]]:
    for row in range(sheet.max_row):
        row_labels = {
            col: _normalize_text(sheet.value(row, col))
            for col in range(sheet.max_col)
        }

        diameter_cols = [
            col for col, text in row_labels.items()
            if "DIAMETRO" in text
        ]
        length_cols = [
            col for col, text in row_labels.items()
            if text == "LARGO" or "LARGO" in text
        ]
        kg_mt_cols = [
            col for col, text in row_labels.items()
            if "KG/MT" in text or "KG/ M" in text
        ]
        bar_cols = [
            col for col, text in row_labels.items()
            if "BARRAS" in text
        ]

        if diameter_cols and length_cols and bar_cols:
            diameter_col = diameter_cols[0]
            length_col = length_cols[0]
            bars_col = bar_cols[0]

            kg_mt_col = kg_mt_cols[0] if kg_mt_cols else length_col + 1

            # En las planillas Bellet, el Kg total está inmediatamente después
            # de Barras. Se busca además un encabezado "KG" independiente.
            total_kg_col = None
            for col, text in row_labels.items():
                if text == "KG" and col != kg_mt_col:
                    total_kg_col = col
                    if col > bars_col:
                        break
            if total_kg_col is None:
                total_kg_col = bars_col + 1

            return row, {
                "diametro": diameter_col,
                "largo": length_col,
                "kg_mt": kg_mt_col,
                "barras": bars_col,
                "kg": total_kg_col,
            }

    raise FierroOrderParseError(
        "No se encontró la cabecera del desglose de fierro."
    )


def _parse_detail(sheet: _SheetMatrix) -> tuple[FierroOrderLine, ...]:
    header_row, columns = _find_detail_header(sheet)
    rows: list[FierroOrderLine] = []
    blank_streak = 0

    # La fila inmediatamente siguiente suele ser la unidad (mm/mts), por eso
    # el parser simplemente descarta filas sin diámetro/largo/barras numéricos.
    for row in range(header_row + 1, sheet.max_row):
        row_text = " ".join(
            _normalize_text(sheet.value(row, col))
            for col in range(sheet.max_col)
            if not _is_blank(sheet.value(row, col))
        )

        if "TOTAL" in row_text and rows:
            break

        diametro = _as_number(sheet.value(row, columns["diametro"]))
        largo = _as_number(sheet.value(row, columns["largo"]))
        barras = _as_int(sheet.value(row, columns["barras"]))

        if diametro is None or largo is None or barras is None:
            if not row_text:
                blank_streak += 1
                if rows and blank_streak >= 4:
                    break
            continue

        blank_streak = 0
        if diametro <= 0 or largo <= 0 or barras <= 0:
            continue

        kg_mt_source = _as_number(sheet.value(row, columns["kg_mt"]))
        kg_mt = kg_mt_source if kg_mt_source and kg_mt_source > 0 else (
            diametro ** 2 / 162.0
        )

        kg_source = _as_number(sheet.value(row, columns["kg"]))
        kg = kg_source if kg_source and kg_source > 0 else (
            barras * largo * kg_mt
        )

        rows.append(
            FierroOrderLine(
                diametro_mm=round(diametro, 3),
                largo_m=round(largo, 3),
                kg_mt=round(kg_mt, 6),
                barras_solicitadas=barras,
                kg_solicitados=round(kg, 6),
            )
        )

    if not rows:
        raise FierroOrderParseError(
            "El pedido no contiene líneas de material reconocibles."
        )

    return tuple(rows)


def _matrix_from_xlsx(content: bytes) -> _SheetMatrix:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FierroOrderParseError(
            "Falta la dependencia openpyxl para leer archivos .xlsx."
        ) from exc

    try:
        workbook = load_workbook(
            BytesIO(content),
            data_only=True,
            read_only=True,
        )
    except Exception as exc:
        raise FierroOrderParseError(
            f"No se pudo abrir el archivo .xlsx: {exc}"
        ) from exc

    sheet = workbook[workbook.sheetnames[0]]
    values = [list(row) for row in sheet.iter_rows(values_only=True)]
    return _SheetMatrix(values=values)


def _matrix_from_xls(content: bytes) -> _SheetMatrix:
    try:
        import xlrd
    except ImportError as exc:
        raise FierroOrderParseError(
            "Falta la dependencia xlrd para leer archivos .xls."
        ) from exc

    try:
        workbook = xlrd.open_workbook(file_contents=content)
        sheet = workbook.sheet_by_index(0)
    except Exception as exc:
        raise FierroOrderParseError(
            f"No se pudo abrir el archivo .xls: {exc}"
        ) from exc

    values: list[list[Any]] = []
    for row in range(sheet.nrows):
        current: list[Any] = []
        for col in range(sheet.ncols):
            cell = sheet.cell(row, col)
            value: Any = cell.value

            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    dt = xlrd.xldate.xldate_as_datetime(
                        cell.value,
                        workbook.datemode,
                    )
                    # Valores menores a 1 representan horas puras.
                    if 0 <= float(cell.value) < 1:
                        value = dt.time()
                    else:
                        value = dt
                except Exception:
                    value = cell.value

            current.append(value)
        values.append(current)

    return _SheetMatrix(values=values)


def parse_fierro_order(
    content: bytes,
    filename: str,
) -> FierroOrder:
    if not content:
        raise FierroOrderParseError("El archivo de pedido está vacío.")

    suffix = PurePosixPath(filename or "").suffix.lower()
    if suffix == ".xls":
        sheet = _matrix_from_xls(content)
    elif suffix == ".xlsx":
        sheet = _matrix_from_xlsx(content)
    else:
        raise FierroOrderParseError(
            f"Formato no soportado para pedido de fierro: {suffix or 'sin extensión'}. "
            "Use .xls o .xlsx."
        )

    raw_number = _find_label_value(
        sheet,
        ("PEDIDO N°", "PEDIDO Nº", "PEDIDO NO", "PEDIDO N"),
        "PEDIDO N°",
    )
    number_as_int = _as_int(raw_number)
    numero_pedido = (
        str(number_as_int)
        if number_as_int is not None
        else str(raw_number).strip()
    )

    area = str(
        _find_label_value(sheet, ("AREA",), "AREA")
    ).strip()

    fecha_pedido = _as_date(
        _find_label_value(
            sheet,
            ("FECHA PEDIDO",),
            "FECHA PEDIDO",
        ),
        "FECHA PEDIDO",
    )

    fecha_obra = _as_date(
        _find_label_value(
            sheet,
            ("FECHA EN OBRA", "FECHA OBRA"),
            "FECHA EN OBRA",
        ),
        "FECHA EN OBRA",
    )

    raw_time = _find_label_value(
        sheet,
        ("HORA EN OBRA", "HORA OBRA"),
        "HORA EN OBRA",
    )
    hora = _as_time_text(raw_time)

    lines = _parse_detail(sheet)

    return FierroOrder(
        numero_pedido=numero_pedido,
        area=area,
        fecha_pedido=fecha_pedido,
        fecha_requerida_obra=fecha_obra,
        hora_requerida=hora,
        lineas=lines,
    )
