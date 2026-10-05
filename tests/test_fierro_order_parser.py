from __future__ import annotations

from datetime import datetime, time
from io import BytesIO

from openpyxl import Workbook

from app.fierro_order_parser import parse_fierro_order


def _build_order_xlsx(
    number: int,
    area: str,
    order_date: datetime,
    required_date: datetime,
    required_time,
    lines: list[tuple[float, float, int]],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Nº{number}"

    ws["B8"] = "PEDIDO N°"
    ws["D8"] = number
    ws["B9"] = "AREA"
    ws["D9"] = area
    ws["B10"] = "FECHA PEDIDO"
    ws["D10"] = order_date
    ws["B11"] = "FECHA EN OBRA"
    ws["D11"] = required_date
    ws["B12"] = "HORA EN OBRA"
    ws["D12"] = required_time

    ws["B15"] = "Diametro"
    ws["C15"] = "Largo"
    ws["D15"] = "Kg/mt"
    ws["E15"] = "Barras"
    ws["F15"] = "Kg"

    row = 17
    for diameter, length, bars in lines:
        kg_mt = diameter ** 2 / 162
        ws.cell(row, 2, diameter)
        ws.cell(row, 3, length)
        ws.cell(row, 4, kg_mt)
        ws.cell(row, 5, bars)
        ws.cell(row, 6, kg_mt * length * bars)
        row += 1

    ws.cell(28, 5, "TOTAL")
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_pedido_4_preserva_precision_del_total():
    content = _build_order_xlsx(
        4,
        "sub -2",
        datetime(2026, 9, 16),
        datetime(2026, 9, 21),
        "10 y 12 am",
        [
            (25, 9, 60),
            (25, 7, 54),
            (25, 12, 18),
            (16, 12, 106),
            (22, 12, 56),
            (10, 12, 270),
        ],
    )

    order = parse_fierro_order(content, "Pedido Fierro Nº4_Bellet.xlsx")

    assert order.numero_pedido == "4"
    assert order.area == "sub -2"
    assert order.fecha_pedido.isoformat() == "2026-09-16"
    assert order.fecha_requerida_obra.isoformat() == "2026-09-21"
    assert order.hora_requerida == "10 y 12 am"
    assert len(order.lineas) == 6
    assert order.total_kg == 10392.78


def test_hora_excel_se_convierte_a_texto():
    content = _build_order_xlsx(
        5,
        "GRUA",
        datetime(2026, 9, 28),
        datetime(2026, 9, 29),
        time(9, 0),
        [(18, 10, 25)],
    )

    order = parse_fierro_order(content, "Pedido Fierro Nº5_Bellet.xlsx")
    assert order.hora_requerida == "09:00"
    assert order.total_kg == 500.0
